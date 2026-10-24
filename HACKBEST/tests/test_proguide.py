import csv
import html
from datetime import date, datetime, timedelta
from io import BytesIO, StringIO
from urllib.parse import parse_qs, urlparse

from app import create_app
from app.email_service import send_password_changed_email, send_password_reset_email
from app.models import (
    Achievement,
    Action,
    ActionAssignment,
    AllocationHistory,
    AuditLog,
    Concern,
    GroupChat,
    GroupChatMember,
    GroupMessage,
    Meeting,
    Message,
    MessageDelivery,
    Notification,
    PasswordResetToken,
    Substitution,
    User,
    UserAuthState,
    db,
)
from openpyxl import Workbook, load_workbook


def build_app(tmp_path):
    db_path = tmp_path / "proguide.db"
    import config

    config.Config.SQLALCHEMY_DATABASE_URI = f"sqlite:///{db_path}"
    app = create_app()
    app.instance_path = str(tmp_path / "instance")
    (tmp_path / "instance").mkdir(exist_ok=True)
    app.config["WTF_CSRF_ENABLED"] = False
    app.config["TESTING"] = True
    return app


def test_admin_is_seeded_and_forces_change_password(tmp_path):
    app = build_app(tmp_path)
    client = app.test_client()
    password = app.config["SEED_ADMIN_PASSWORD"]
    resp = client.post(
        "/login",
        data={"username": "admin", "password": password, "role": "admin"},
        follow_redirects=True,
    )
    assert resp.status_code == 200
    assert b"Update password" in resp.data


def test_seeded_admin_uses_configured_recovery_email(tmp_path, monkeypatch):
    import config

    monkeypatch.setattr(config.Config, "SQLALCHEMY_DATABASE_URI", f"sqlite:///{tmp_path / 'proguide.db'}")
    monkeypatch.setattr(config.Config, "SEED_ADMIN_EMAIL", "admin@example.test", raising=False)
    app = create_app()
    with app.app_context():
        assert User.query.filter_by(username="admin").one().email == "admin@example.test"
        db.session.remove()
        db.engine.dispose()


def test_allocation_table_shows_current_mentor_name(tmp_path):
    app = build_app(tmp_path)
    client = app.test_client()
    with app.app_context():
        admin = User.query.filter_by(username="admin").one()
        admin.must_change_password = False
        mentor = User(
            username="allocation_mentor",
            full_name="Jordan Mentor",
            role="mentor",
            is_active=False,
        )
        mentee = User(
            username="allocation_mentee",
            full_name="Casey Mentee",
            role="mentee",
            is_active=True,
        )
        mentor.set_password("MentorTest!2026")
        mentee.set_password("MenteeTest!2026")
        db.session.add_all([mentor, mentee])
        db.session.commit()
        mentee.mentor_id = mentor.id
        unassigned = User(
            username="unassigned_mentee",
            full_name="Avery Unassigned",
            role="mentee",
            is_active=True,
        )
        unassigned.set_password("MenteeTest!2026")
        db.session.add(unassigned)
        db.session.commit()
        admin_password = app.config["SEED_ADMIN_PASSWORD"]

    client.post(
        "/login",
        data={"username": "admin", "password": admin_password, "role": "admin"},
    )
    response = client.get("/admin/allocations")
    assert response.status_code == 200
    assert b"Jordan Mentor" in response.data
    assert b"Casey Mentee" in response.data
    assert b"Unassigned" in response.data
    assert b">2<" not in response.data
    with app.app_context():
        db.session.remove()
        db.engine.dispose()


def test_admin_can_change_password_and_access_dashboard(tmp_path):
    app = build_app(tmp_path)
    client = app.test_client()
    password = app.config["SEED_ADMIN_PASSWORD"]
    client.post(
        "/login",
        data={"username": "admin", "password": password, "role": "admin"},
        follow_redirects=True,
    )
    new_password = "NewStrongPassword!1"
    resp = client.post(
        "/change-password",
        data={
            "current_password": password,
            "password": new_password,
            "confirm_password": new_password,
        },
        follow_redirects=True,
    )
    assert resp.status_code == 200
    assert b"Password updated successfully" in resp.data

    client.get("/logout", follow_redirects=True)
    login_resp = client.post(
        "/login",
        data={"username": "admin", "password": new_password, "role": "admin"},
        follow_redirects=True,
    )
    assert login_resp.status_code == 200
    assert b"Admin dashboard" in login_resp.data


def test_admin_can_create_user_without_optional_email(tmp_path):
    app = build_app(tmp_path)
    client = app.test_client()
    client.post(
        "/login",
        data={"username": "admin", "password": app.config["SEED_ADMIN_PASSWORD"], "role": "admin"},
        follow_redirects=True,
    )
    client.post(
        "/change-password",
        data={
            "current_password": app.config["SEED_ADMIN_PASSWORD"],
            "password": "DemoAdminPass!2026",
            "confirm_password": "DemoAdminPass!2026",
        },
        follow_redirects=True,
    )

    response = client.post(
        "/admin/users/new",
        data={
            "username": "mentor_demo",
            "full_name": "Morgan Mentor",
            "role": "mentor",
            "password": "MentorTemp!2026",
            "submit": "Save",
        },
        follow_redirects=True,
    )

    assert response.status_code == 200
    assert b"mentor_demo" in response.data
    assert b'value="MentorTemp!2026"' in response.data
    assert b"MentorTemp!2026" not in response.headers.get("Set-Cookie", "").encode()
    assert "no-store" in response.headers["Cache-Control"]
    with app.app_context():
        user = User.query.filter_by(username="mentor_demo").first()
        assert user is not None
        assert user.email is None
        db.session.remove()
        db.engine.dispose()


def test_password_change_requires_current_password_policy_and_revokes_other_sessions(tmp_path):
    app = build_app(tmp_path)
    client = app.test_client()
    other_session = app.test_client()
    with app.app_context():
        user = User(
            username="secure_change",
            full_name="Secure Change",
            role="mentee",
            must_change_password=False,
        )
        user.set_password("OriginalPassword!1")
        db.session.add(user)
        db.session.commit()
        user_id = user.id

    login_data = {"username": "secure_change", "password": "OriginalPassword!1", "role": "mentee"}
    client.post("/login", data=login_data)
    other_session.post("/login", data=login_data)

    wrong_current = client.post(
        "/change-password",
        data={
            "current_password": "WrongPassword!1",
            "password": "FreshPassword!2026",
            "confirm_password": "FreshPassword!2026",
        },
    )
    assert wrong_current.status_code == 400
    assert b"The current password is incorrect." in wrong_current.data

    weak_password = client.post(
        "/change-password",
        data={
            "current_password": "OriginalPassword!1",
            "password": "weakpassword",
            "confirm_password": "weakpassword",
        },
    )
    assert weak_password.status_code == 400
    assert b"uppercase, lowercase, a number, and a special character" in weak_password.data

    long_password = "A" * 70 + "a1!?"
    overlong = client.post(
        "/change-password",
        data={
            "current_password": "OriginalPassword!1",
            "password": long_password,
            "confirm_password": long_password,
        },
    )
    assert overlong.status_code == 400
    assert b"at most 72 UTF-8 bytes" in overlong.data

    changed = client.post(
        "/change-password",
        data={
            "current_password": "OriginalPassword!1",
            "password": "FreshPassword!2026",
            "confirm_password": "FreshPassword!2026",
        },
        follow_redirects=True,
    )
    assert changed.status_code == 200
    assert b"Password updated successfully" in changed.data
    assert client.get("/mentee/dashboard").status_code == 200
    assert other_session.get("/mentee/dashboard").status_code == 302
    with app.app_context():
        user = db.session.get(User, user_id)
        assert user.check_password("FreshPassword!2026")
        assert user.must_change_password is False
        assert db.session.get(UserAuthState, user_id).session_version == 2
        db.session.remove()
        db.engine.dispose()


def test_temporary_password_is_enforced_until_changed_and_seed_file_is_removed(tmp_path):
    app = build_app(tmp_path)
    client = app.test_client()
    seed_file = tmp_path / "instance" / "admin_seed.txt"
    seed_file.write_text("one-time-bootstrap-secret")

    client.post(
        "/login",
        data={
            "username": "admin",
            "password": app.config["SEED_ADMIN_PASSWORD"],
            "role": "admin",
        },
    )
    assert not seed_file.exists()
    assert client.get("/admin/").status_code == 302
    assert client.get("/profile").status_code == 302
    assert client.get("/change-password").status_code == 200

    changed = client.post(
        "/change-password",
        data={
            "current_password": app.config["SEED_ADMIN_PASSWORD"],
            "password": "FreshAdminPassword!2026",
            "confirm_password": "FreshAdminPassword!2026",
        },
    )
    assert changed.status_code == 302
    assert changed.headers["Location"] == "/"
    assert client.get("/admin/").status_code == 200
    with app.app_context():
        db.session.remove()
        db.engine.dispose()


def test_admin_deactivation_and_password_reset_revoke_target_sessions(tmp_path):
    app = build_app(tmp_path)
    admin_client = app.test_client()
    user_client = app.test_client()
    with app.app_context():
        admin = User.query.filter_by(username="admin").one()
        admin.must_change_password = False
        user = User(
            username="revoke_target",
            full_name="Revoke Target",
            role="mentor",
            must_change_password=False,
        )
        user.set_password("TargetPassword!2026")
        db.session.add(user)
        db.session.commit()
        admin_password = app.config["SEED_ADMIN_PASSWORD"]
        user_id = user.id

    admin_client.post(
        "/login",
        data={"username": "admin", "password": admin_password, "role": "admin"},
    )
    user_login_data = {"username": "revoke_target", "password": "TargetPassword!2026", "role": "mentor"}
    user_client.post("/login", data=user_login_data)

    deactivated = admin_client.post(f"/admin/users/{user_id}/toggle")
    assert deactivated.status_code == 302
    assert user_client.get("/mentor/dashboard").status_code == 302

    admin_client.post(f"/admin/users/{user_id}/toggle")
    user_client.post("/login", data=user_login_data)
    reset_response = admin_client.post(f"/admin/users/{user_id}/reset-password")
    assert reset_response.status_code == 200
    assert b"One-time temporary password" in reset_response.data
    assert "no-store" in reset_response.headers["Cache-Control"]
    temporary_password = (
        reset_response.data.split(b'id="one-time-password" type="text" value="', 1)[1]
        .split(b'"', 1)[0]
        .decode()
    )
    temporary_password = html.unescape(temporary_password)
    assert temporary_password not in reset_response.headers.get("Set-Cookie", "")
    assert user_client.get("/mentor/dashboard").status_code == 302
    with app.app_context():
        user = db.session.get(User, user_id)
        assert user.check_password(temporary_password)
        assert user.must_change_password is True
        db.session.remove()
        db.engine.dispose()


def test_admin_can_delete_user_and_related_records(tmp_path):
    app = build_app(tmp_path)
    client = app.test_client()
    with app.app_context():
        admin = User.query.filter_by(username="admin").one()
        admin.must_change_password = False
        db.session.commit()
        admin_id = admin.id
        mentor = User(
            username="delete_mentor",
            full_name="Morgan Delete",
            role="mentor",
            is_active=True,
        )
        mentee = User(
            username="delete_mentee",
            full_name="Taylor Keep",
            role="mentee",
            is_active=True,
        )
        mentor.set_password("DeleteTest!2026")
        mentee.set_password("DeleteTest!2026")
        db.session.add_all([mentor, mentee])
        db.session.commit()
        mentee.mentor_id = mentor.id
        action = Action(
            title="Delete-owned action",
            description="Action created by deleted mentor",
            deadline=date.today() + timedelta(days=4),
            created_by_id=mentor.id,
        )
        db.session.add(action)
        db.session.commit()
        assignment = ActionAssignment(action_id=action.id, mentee_id=mentee.id)
        meeting = Meeting(
            mentee_id=mentee.id,
            mentor_id=mentor.id,
            requested_by_id=mentor.id,
            status="Approved",
            date=date.today() + timedelta(days=2),
            time="10:00",
            duration=60,
            mode="online",
            agenda="Delete this meeting",
        )
        concern = Concern(
            mentee_id=mentee.id,
            mentor_id=mentor.id,
            title="Keep mentee concern",
            category="academic",
            description="This concern belongs to the mentee and should remain.",
        )
        message = Message(sender_id=mentor.id, recipient_id=mentee.id, content="Delete this message")
        db.session.add_all([assignment, meeting, concern, message])
        db.session.commit()
        delivery = MessageDelivery(message_id=message.id)
        notification = Notification(
            user_id=mentor.id,
            type="test",
            message="Delete this notification",
        )
        reset_token = PasswordResetToken(
            user_id=mentor.id,
            token_hash="a" * 64,
            expires_at=datetime.now() + timedelta(minutes=15),
        )
        auth_state = UserAuthState(user_id=mentor.id, session_version=2)
        history = AllocationHistory(
            mentee_id=mentee.id,
            old_mentor_id=None,
            new_mentor_id=mentor.id,
            changed_by_id=admin.id,
        )
        db.session.add_all([delivery, notification, reset_token, auth_state, history])
        db.session.commit()
        mentor_id, mentee_id = mentor.id, mentee.id

    client.post(
        "/login",
        data={"username": "admin", "password": app.config["SEED_ADMIN_PASSWORD"], "role": "admin"},
    )
    users_page = client.get("/admin/users")
    assert users_page.status_code == 200
    assert users_page.data.count(b"Delete account") == 2
    assert b'data-confirm-delete="System Administrator"' not in users_page.data
    deleted = client.post(f"/admin/users/{mentor_id}/delete", follow_redirects=True)
    assert deleted.status_code == 200
    assert b"Deleted Morgan Delete and their related records." in deleted.data
    with app.app_context():
        assert db.session.get(User, mentor_id) is None
        kept_mentee = db.session.get(User, mentee_id)
        assert kept_mentee is not None
        assert kept_mentee.mentor_id is None
        assert Meeting.query.count() == 0
        assert Message.query.count() == 0
        assert MessageDelivery.query.count() == 0
        assert Action.query.count() == 0
        assert ActionAssignment.query.count() == 0
        kept_concern = Concern.query.one()
        assert kept_concern.mentee_id == mentee_id
        assert kept_concern.mentor_id is None
        assert Notification.query.filter_by(user_id=mentor_id).count() == 0
        assert PasswordResetToken.query.filter_by(user_id=mentor_id).count() == 0
        assert UserAuthState.query.filter_by(user_id=mentor_id).count() == 0
        assert AllocationHistory.query.count() == 0
        assert AuditLog.query.filter_by(user_id=admin_id, action="user_delete").count() == 1
        db.session.remove()
        db.engine.dispose()


def test_admin_cannot_delete_self(tmp_path):
    app = build_app(tmp_path)
    client = app.test_client()
    with app.app_context():
        admin = User.query.filter_by(username="admin").one()
        admin.must_change_password = False
        db.session.commit()
        admin_id = admin.id
    client.post(
        "/login",
        data={"username": "admin", "password": app.config["SEED_ADMIN_PASSWORD"], "role": "admin"},
    )

    self_delete = client.post(f"/admin/users/{admin_id}/delete", follow_redirects=True)
    assert self_delete.status_code == 200
    assert b"You cannot delete the account you are currently using." in self_delete.data
    with app.app_context():
        assert db.session.get(User, admin_id) is not None
        db.session.remove()
        db.engine.dispose()


def test_role_dashboards_render_scoped_navigation_and_access_denials(tmp_path):
    app = build_app(tmp_path)
    app.config["WTF_CSRF_ENABLED"] = False
    client = app.test_client()
    with app.app_context():
        admin = User.query.filter_by(username="admin").first()
        admin.must_change_password = False
        mentor = User(username="mentor_ui", full_name="Morgan Mentor", role="mentor", is_active=True)
        substitute = User(username="substitute_ui", full_name="Riley Substitute", role="mentor", is_active=True)
        mentee = User(username="mentee_ui", full_name="Taylor Mentee", role="mentee", is_active=True)
        mentor.set_password("MentorDemo!2026")
        substitute.set_password("SubstituteDemo!2026")
        mentee.set_password("MenteeDemo!2026")
        db.session.add_all([mentor, substitute, mentee])
        db.session.commit()
        mentee.mentor_id = mentor.id
        meeting = Meeting(
            mentee_id=mentee.id,
            mentor_id=mentor.id,
            requested_by_id=mentee.id,
            status="Approved",
            date=date.today() + timedelta(days=1),
            time="10:00",
            duration=60,
            mode="online",
            agenda="UI access test",
        )
        substitution = Substitution(
            mentee_id=mentee.id,
            mentor_id=mentor.id,
            substitute_mentor_id=substitute.id,
            reason="Coverage",
            start_date=date.today(),
            end_date=date.today() + timedelta(days=2),
            active=True,
        )
        db.session.add_all([meeting, substitution])
        db.session.commit()

    login = client.post(
        "/login",
        data={"username": "admin", "password": app.config["SEED_ADMIN_PASSWORD"], "role": "admin"},
        follow_redirects=True,
    )
    assert login.status_code == 200
    admin_page = client.get("/admin/")
    assert admin_page.status_code == 200
    assert b"User management" in admin_page.data
    assert b"My mentees" not in admin_page.data

    client.get("/logout")
    client.post("/login", data={"username": "mentor_ui", "password": "MentorDemo!2026", "role": "mentor"})
    mentor_page = client.get("/mentor/dashboard")
    assert mentor_page.status_code == 200
    assert b"My mentees" in mentor_page.data
    assert b"User management" not in mentor_page.data
    assert client.get("/admin/").status_code == 403
    assert b"ACCESS RESTRICTED" in client.get("/admin/").data
    search = client.get("/api/search?q=Taylor")
    assert search.status_code == 200
    assert [person["name"] for person in search.get_json()] == ["Taylor Mentee"]

    client.get("/logout")
    client.post("/login", data={"username": "mentee_ui", "password": "MenteeDemo!2026", "role": "mentee"})
    mentee_page = client.get("/mentee/dashboard")
    assert mentee_page.status_code == 200
    assert b"Book a session" in mentee_page.data
    assert b"Reports & analytics" not in mentee_page.data
    assert client.get("/mentor/dashboard").status_code == 403
    client.get("/logout")
    client.post("/login", data={"username": "substitute_ui", "password": "SubstituteDemo!2026", "role": "mentor"})
    assert client.get("/mentor/meetings/1").status_code == 200
    with app.app_context():
        db.session.remove()
        db.engine.dispose()


def test_role_selection_login_enforces_role_and_lockout(tmp_path):
    app = build_app(tmp_path)
    client = app.test_client()
    with app.app_context():
        admin = User.query.filter_by(username="admin").first()
        admin.must_change_password = False
        db.session.commit()

    selection = client.get("/login")
    assert b"Choose your workspace" in selection.data
    assert all(label in selection.data for label in (b"Admin", b"Mentor", b"Mentee"))

    missing_role = client.post(
        "/login",
        data={"username": "admin", "password": app.config["SEED_ADMIN_PASSWORD"]},
    )
    assert b"Please choose a role before signing in." in missing_role.data

    mismatch = client.post(
        "/login",
        data={"username": "admin", "password": app.config["SEED_ADMIN_PASSWORD"], "role": "mentor"},
    )
    assert b"We couldn&#39;t sign you in with those details" in mismatch.data

    for _ in range(4):
        response = client.post(
            "/login",
            data={"username": "admin", "password": "WrongPassword!1", "role": "admin"},
        )
        assert b"We couldn&#39;t sign you in with those details" in response.data
    locked = client.post(
        "/login",
        data={"username": "admin", "password": app.config["SEED_ADMIN_PASSWORD"], "role": "admin"},
    )
    assert b"Too many failed attempts" in locked.data
    with app.app_context():
        db.session.remove()
        db.engine.dispose()


def test_forgot_password_and_single_use_reset_flow(tmp_path, monkeypatch):
    app = build_app(tmp_path)
    app.config["FRONTEND_URL"] = "https://proguide.example"
    client = app.test_client()
    with app.app_context():
        user = User(username="reset_user", full_name="Reset User", email="reset@example.com", role="mentor")
        user.set_password("OriginalPassword!1")
        db.session.add(user)
        db.session.commit()
        user_id = user.id

    sent_links = []
    sent_confirmations = []

    def capture_reset_email(user, reset_url):
        sent_links.append(reset_url)

    def capture_confirmation_email(user):
        sent_confirmations.append(user.id)

    monkeypatch.setattr("app.blueprints.auth.send_password_reset_email", capture_reset_email)
    monkeypatch.setattr("app.blueprints.auth.send_password_changed_email", capture_confirmation_email)
    neutral = b"If the details match an account, a password reset link has been sent to the registered email."
    mismatch = client.post(
        "/forgot-password",
        data={"username": "reset_user", "email": "wrong@example.com", "role": "mentor"},
        follow_redirects=True,
    )
    assert neutral in mismatch.data
    assert not sent_links

    matching = client.post(
        "/forgot-password",
        data={"username": "reset_user", "email": "reset@example.com", "role": "mentor"},
        follow_redirects=True,
    )
    assert neutral in matching.data
    assert len(sent_links) == 1
    assert sent_links[0].startswith("https://proguide.example/reset-password?token=")
    token_value = parse_qs(urlparse(sent_links[0]).query)["token"][0]
    legacy_session_client = app.test_client()
    with legacy_session_client.session_transaction() as legacy_session:
        legacy_session["_user_id"] = str(user_id)
        legacy_session["_fresh"] = True
    existing_login_client = app.test_client()
    existing_login_client.post(
        "/login",
        data={"username": "reset_user", "password": "OriginalPassword!1", "role": "mentor"},
    )
    assert legacy_session_client.get("/profile").status_code == 200
    with app.app_context():
        stored_token = PasswordResetToken.query.filter_by(user_id=user_id).one()
        assert token_value not in stored_token.token_hash
        assert stored_token.expires_at > stored_token.created_at
        assert round((stored_token.expires_at - stored_token.created_at).total_seconds()) == 900

    invalid_password = client.post(
        "/reset-password",
        data={
            "token": token_value,
            "password": "weakpassword",
            "confirm_password": "weakpassword",
        },
    )
    assert b"uppercase, lowercase, a number, and a special character" in invalid_password.data

    same_password = client.post(
        "/reset-password",
        data={
            "token": token_value,
            "password": "OriginalPassword!1",
            "confirm_password": "OriginalPassword!1",
        },
    )
    assert b"Choose a password different from your current password" in same_password.data

    reset = client.post(
        "/reset-password",
        data={
            "token": token_value,
            "password": "FreshPassword!2026",
            "confirm_password": "FreshPassword!2026",
        },
    )
    assert reset.status_code == 200
    assert b"Password updated" in reset.data
    assert sent_confirmations == [user_id]
    assert client.get(f"/reset-password?token={token_value}").status_code == 400
    with app.app_context():
        updated_user = db.session.get(User, user_id)
        assert updated_user.check_password("FreshPassword!2026")
        assert not updated_user.check_password("OriginalPassword!1")
        assert db.session.get(UserAuthState, user_id).session_version == 2
    assert existing_login_client.get("/profile").status_code == 302
    assert legacy_session_client.get("/profile").status_code == 302
    with app.app_context():
        db.session.remove()
        db.engine.dispose()


def test_expired_reset_token_is_rejected(tmp_path):
    app = build_app(tmp_path)
    client = app.test_client()
    with app.app_context():
        user = User(username="expired_user", full_name="Expired User", email="expired@example.com", role="mentee")
        user.set_password("OriginalPassword!1")
        db.session.add(user)
        db.session.flush()
        token = PasswordResetToken(
            user_id=user.id,
            token_hash="a" * 64,
            expires_at=date.today(),
            used=False,
        )
        db.session.add(token)
        db.session.commit()

    response = client.get(f"/reset-password?token={'a' * 64}")
    assert response.status_code == 400
    assert b"invalid, expired, or has already been used" in response.data
    with app.app_context():
        db.session.remove()
        db.engine.dispose()


def test_forgot_password_requests_are_rate_limited(tmp_path, monkeypatch):
    app = build_app(tmp_path)
    client = app.test_client()
    with app.app_context():
        user = User(username="limited_user", full_name="Limited User", email="limited@example.com", role="mentor")
        user.set_password("OriginalPassword!1")
        db.session.add(user)
        db.session.commit()

    sent_links = []
    monkeypatch.setattr(
        "app.blueprints.auth.send_password_reset_email",
        lambda user, reset_url: sent_links.append(reset_url),
    )
    data = {"username": "limited_user", "email": "limited@example.com", "role": "mentor"}
    for _ in range(6):
        response = client.post("/forgot-password", data=data, follow_redirects=True)
        assert b"If the details match an account" in response.data
    assert len(sent_links) == 5
    with app.app_context():
        db.session.remove()
        db.engine.dispose()


def test_password_email_templates_render_expected_content(tmp_path, monkeypatch):
    app = build_app(tmp_path)
    sent_messages = []
    monkeypatch.setattr(
        "app.email_service._send_email",
        lambda recipient, subject, text, html: sent_messages.append((recipient, subject, text, html)),
    )
    with app.test_request_context("/"):
        user = User(
            username="mail_user",
            full_name="<Mentee>",
            email="mail@example.com",
            role="mentee",
            password_hash="unused",
        )
        reset_url = "https://proguide.example/reset-password?token=one-time-token"
        send_password_reset_email(user, reset_url)
        send_password_changed_email(user)

    assert len(sent_messages) == 2
    assert sent_messages[0][0] == "mail@example.com"
    assert "Reset Password" in sent_messages[0][3]
    assert reset_url in sent_messages[0][2]
    assert "15 minutes" in sent_messages[0][3]
    assert reset_url in sent_messages[0][3]
    assert "If you did not request this, ignore this email." in sent_messages[0][3]
    assert "&lt;Mentee&gt;" in sent_messages[0][3]
    assert "Your ProGuide password was changed." in sent_messages[1][2]
    with app.app_context():
        db.session.remove()
        db.engine.dispose()


def test_mentee_inbox_only_offers_and_accepts_mentor_recipients(tmp_path):
    app = build_app(tmp_path)
    client = app.test_client()
    with app.app_context():
        admin = User.query.filter_by(username="admin").one()
        admin.must_change_password = False
        mentor = User(username="message_mentor", full_name="Morgan Mentor", role="mentor", is_active=True)
        unrelated_mentor = User(
            username="unrelated_mentor",
            full_name="Unrelated Mentor",
            role="mentor",
            is_active=True,
        )
        inactive_mentor = User(
            username="inactive_mentor",
            full_name="Inactive Mentor",
            role="mentor",
            is_active=False,
        )
        other_mentee = User(username="other_mentee", full_name="Other Mentee", role="mentee", is_active=True)
        mentee = User(username="message_mentee", full_name="Taylor Mentee", role="mentee", is_active=True)
        for user in (mentor, unrelated_mentor, inactive_mentor, other_mentee, mentee):
            user.set_password("MessageTest!2026")
        db.session.add_all([mentor, unrelated_mentor, inactive_mentor, other_mentee, mentee])
        db.session.commit()
        mentee.mentor_id = mentor.id
        db.session.commit()
        mentor_id = mentor.id
        admin_id = admin.id
        other_mentee_id = other_mentee.id
        unrelated_mentor_id = unrelated_mentor.id

    client.post(
        "/login",
        data={"username": "message_mentee", "password": "MessageTest!2026", "role": "mentee"},
    )
    inbox = client.get("/messages/")
    assert inbox.status_code == 200
    assert b"Morgan Mentor (mentor)" in inbox.data
    assert b"Unrelated Mentor" not in inbox.data
    assert b"Inactive Mentor" not in inbox.data
    assert b"System Administrator (admin)" not in inbox.data
    assert b"Other Mentee" not in inbox.data

    valid_message = client.post(
        "/messages/",
        data={"recipient_id": str(mentor_id), "content": "Hello, mentor"},
        follow_redirects=True,
    )
    assert b"Message sent." in valid_message.data
    for forbidden_id in (admin_id, other_mentee_id, unrelated_mentor_id):
        forbidden_message = client.post(
            "/messages/",
            data={"recipient_id": str(forbidden_id), "content": "This should not be sent"},
            follow_redirects=True,
        )
        assert b"This recipient is not available for your account." in forbidden_message.data

    with app.app_context():
        messages = Message.query.all()
        assert len(messages) == 1
        assert messages[0].recipient_id == mentor_id
        db.session.remove()
        db.engine.dispose()


def test_chat_delivery_seen_receipts_and_recipient_access(tmp_path):
    app = build_app(tmp_path)
    mentee_client = app.test_client()
    mentor_client = app.test_client()
    with app.app_context():
        mentor = User(username="chat_mentor", full_name="Jamie Mentor", role="mentor", is_active=True)
        mentee = User(username="chat_mentee", full_name="Riley Mentee", role="mentee", is_active=True)
        other_mentee = User(username="chat_other", full_name="Other Mentee", role="mentee", is_active=True)
        for user in (mentor, mentee, other_mentee):
            user.set_password("ChatTest!2026")
        db.session.add_all([mentor, mentee, other_mentee])
        db.session.commit()
        mentee.mentor_id = mentor.id
        other_mentee.mentor_id = mentor.id
        db.session.commit()
        mentor_id, mentee_id, other_mentee_id = mentor.id, mentee.id, other_mentee.id

    mentee_client.post(
        "/login",
        data={"username": "chat_mentee", "password": "ChatTest!2026", "role": "mentee"},
    )
    mentor_client.post(
        "/login",
        data={"username": "chat_mentor", "password": "ChatTest!2026", "role": "mentor"},
    )

    sent = mentee_client.post(
        f"/messages/api/thread/{mentor_id}",
        json={"content": "Hello Jamie"},
    )
    assert sent.status_code == 201
    message_id = sent.json["id"]
    assert sent.json["delivered"] is False
    assert mentee_client.get(f"/messages/api/thread/{other_mentee_id}").status_code == 404

    received = mentor_client.get(f"/messages/api/thread/{mentee_id}")
    assert received.status_code == 200
    assert received.json["messages"][0]["content"] == "Hello Jamie"
    with app.app_context():
        assert MessageDelivery.query.filter_by(message_id=message_id).one()
        db.session.remove()

    seen = mentor_client.post(
        f"/messages/api/thread/{mentee_id}/seen",
        json={"message_ids": [message_id]},
    )
    assert seen.status_code == 200
    sender_view = mentee_client.get(f"/messages/api/thread/{mentor_id}")
    status = sender_view.json["messages"][0]
    assert status["delivered"] is True
    assert status["is_read"] is True

    with app.app_context():
        db.session.remove()
        db.engine.dispose()


def test_mentor_group_chat_is_limited_to_active_direct_mentees(tmp_path):
    app = build_app(tmp_path)
    mentor_client = app.test_client()
    mentee_client = app.test_client()
    unrelated_client = app.test_client()
    with app.app_context():
        mentor = User(username="group_mentor", full_name="Morgan Mentor", role="mentor", is_active=True)
        other_mentor = User(username="group_other_mentor", full_name="Jordan Mentor", role="mentor", is_active=True)
        mentee = User(username="group_mentee", full_name="Taylor Mentee", role="mentee", is_active=True)
        second_mentee = User(username="group_second", full_name="Riley Mentee", role="mentee", is_active=True)
        inactive_mentee = User(username="group_inactive", full_name="Inactive Mentee", role="mentee", is_active=False)
        substitute_only = User(username="group_substitute", full_name="Substitute Mentee", role="mentee", is_active=True)
        for user in (mentor, other_mentor, mentee, second_mentee, inactive_mentee, substitute_only):
            user.set_password("GroupTest!2026")
        db.session.add_all([mentor, other_mentor, mentee, second_mentee, inactive_mentee, substitute_only])
        db.session.commit()
        mentee.mentor_id = mentor.id
        second_mentee.mentor_id = mentor.id
        inactive_mentee.mentor_id = mentor.id
        substitute_only.mentor_id = other_mentor.id
        db.session.add(Substitution(
            mentee_id=substitute_only.id,
            mentor_id=other_mentor.id,
            substitute_mentor_id=mentor.id,
            reason="Temporary coverage",
            start_date=date.today(),
            end_date=date.today() + timedelta(days=2),
            active=True,
        ))
        db.session.commit()
        mentor_id, mentee_id = mentor.id, mentee.id
        second_mentee_id = second_mentee.id

    mentor_client.post(
        "/login",
        data={"username": "group_mentor", "password": "GroupTest!2026", "role": "mentor"},
    )
    mentee_client.post(
        "/login",
        data={"username": "group_mentee", "password": "GroupTest!2026", "role": "mentee"},
    )
    unrelated_client.post(
        "/login",
        data={"username": "group_substitute", "password": "GroupTest!2026", "role": "mentee"},
    )

    mentee_inbox = mentee_client.get("/messages/")
    assert mentee_inbox.status_code == 200
    assert b"Morgan Mentor&#39;s Mentee Group" in mentee_inbox.data
    mentor_inbox = mentor_client.get("/messages/")
    assert mentor_inbox.status_code == 200
    assert b"Morgan Mentor&#39;s Mentee Group" in mentor_inbox.data
    with app.app_context():
        group = GroupChat.query.filter_by(mentor_id=mentor_id).one()
        group_id = group.id
        member_ids = {
            member.user_id
            for member in GroupChatMember.query.filter_by(group_id=group_id).all()
        }
        assert member_ids == {mentor_id, mentee_id, second_mentee_id}
        db.session.remove()

    assert mentee_client.get(f"/messages/api/group/{group_id}").status_code == 200
    assert unrelated_client.get(f"/messages/api/group/{group_id}").status_code == 404
    sent = mentee_client.post(
        f"/messages/api/group/{group_id}",
        json={"content": "Hello, group"},
    )
    assert sent.status_code == 201
    assert sent.json["sender_name"] == "Taylor Mentee"
    with app.app_context():
        assert GroupMessage.query.filter_by(group_id=group_id, content="Hello, group").count() == 1
        assert Notification.query.filter_by(
            type="group_message",
            related_id=group_id,
        ).count() == 2
        db.session.remove()
        db.engine.dispose()


def test_mentee_can_upload_achievement_and_mentor_is_notified_and_scoped(tmp_path):
    app = build_app(tmp_path)
    mentee_client = app.test_client()
    mentor_client = app.test_client()
    other_mentor_client = app.test_client()
    other_mentee_client = app.test_client()
    with app.app_context():
        mentor = User(username="achievement_mentor", full_name="Morgan Mentor", role="mentor", is_active=True)
        other_mentor = User(
            username="achievement_other_mentor",
            full_name="Jordan Mentor",
            role="mentor",
            is_active=True,
        )
        mentee = User(username="achievement_mentee", full_name="Taylor Mentee", role="mentee", is_active=True)
        other_mentee = User(
            username="achievement_other_mentee",
            full_name="Riley Mentee",
            role="mentee",
            is_active=True,
        )
        for user in (mentor, other_mentor, mentee, other_mentee):
            user.set_password("AchievementTest!2026")
        db.session.add_all([mentor, other_mentor, mentee, other_mentee])
        db.session.commit()
        mentee.mentor_id = mentor.id
        other_mentee.mentor_id = other_mentor.id
        db.session.commit()
        mentee_id, other_mentee_id = mentee.id, other_mentee.id
        mentor_id, other_mentor_id = mentor.id, other_mentor.id

    for client, username, role in (
        (mentee_client, "achievement_mentee", "mentee"),
        (other_mentee_client, "achievement_other_mentee", "mentee"),
        (mentor_client, "achievement_mentor", "mentor"),
        (other_mentor_client, "achievement_other_mentor", "mentor"),
    ):
        client.post(
            "/login",
            data={"username": username, "password": "AchievementTest!2026", "role": role},
        )

    uploaded = mentee_client.post(
        "/mentee/achievements",
        data={
            "title": "Completed robotics project",
            "proof": (BytesIO(b"%PDF-1.7 achievement proof"), "robotics.pdf"),
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert uploaded.status_code == 200
    assert b"Achievement shared with your mentor." in uploaded.data
    assert b"Completed robotics project" in uploaded.data
    with app.app_context():
        achievement = Achievement.query.filter_by(mentee_id=mentee_id).one()
        achievement_id = achievement.id
        assert Notification.query.filter_by(
            user_id=mentor_id,
            type="achievement",
            related_id=achievement_id,
        ).count() == 1
        assert Notification.query.filter_by(
            user_id=other_mentor_id,
            type="achievement",
        ).count() == 0
        db.session.remove()

    mentor_notifications = mentor_client.get("/api/notifications")
    assert mentor_notifications.status_code == 200
    notification = next(
        item for item in mentor_notifications.json if item["type"] == "achievement"
    )
    assert notification["url"] == f"/mentor/achievements/{mentee_id}"

    own_proof = mentee_client.get(f"/mentee/achievements/{achievement_id}/proof")
    assert own_proof.status_code == 200
    assert own_proof.data == b"%PDF-1.7 achievement proof"
    assert other_mentee_client.get(f"/mentee/achievements/{achievement_id}/proof").status_code == 404

    mentor_page = mentor_client.get("/mentor/achievements")
    assert mentor_page.status_code == 200
    assert b"Students Achievements" in mentor_page.data
    assert b"Taylor Mentee" in mentor_page.data
    assert b"Completed robotics project" in mentor_page.data
    individual_page = mentor_client.get(f"/mentor/achievements/{mentee_id}")
    assert individual_page.status_code == 200
    assert b"Completed robotics project" in individual_page.data
    assert not any(
        item["type"] == "achievement"
        for item in mentor_client.get("/api/notifications").json
    )
    mentor_proof = mentor_client.get(
        f"/mentor/achievements/{mentee_id}/{achievement_id}/proof"
    )
    assert mentor_proof.status_code == 200
    assert mentor_proof.data == b"%PDF-1.7 achievement proof"
    assert other_mentor_client.get(f"/mentor/achievements/{mentee_id}").status_code == 404
    assert other_mentor_client.get(
        f"/mentor/achievements/{mentee_id}/{achievement_id}/proof"
    ).status_code == 404

    bad_proof = mentee_client.post(
        "/mentee/achievements",
        data={
            "title": "Bad file",
            "proof": (BytesIO(b"not a pdf"), "fake.pdf"),
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert b"does not match its file type" in bad_proof.data
    with app.app_context():
        assert Achievement.query.filter_by(mentee_id=mentee_id).count() == 1
        db.session.remove()
        db.engine.dispose()


def test_mentor_can_create_meeting_for_assigned_mentee(tmp_path):
    app = build_app(tmp_path)
    client = app.test_client()
    with app.app_context():
        mentor = User(username="schedule_mentor", full_name="Morgan Mentor", role="mentor", is_active=True)
        mentee = User(username="schedule_mentee", full_name="Taylor Mentee", role="mentee", is_active=True)
        second_mentee = User(username="schedule_second", full_name="Jordan Mentee", role="mentee", is_active=True)
        third_mentee = User(username="schedule_third", full_name="Casey Mentee", role="mentee", is_active=True)
        other_mentee = User(username="schedule_other", full_name="Other Mentee", role="mentee", is_active=True)
        for user in (mentor, mentee, second_mentee, third_mentee, other_mentee):
            user.set_password("ScheduleTest!2026")
        db.session.add_all([mentor, mentee, second_mentee, third_mentee, other_mentee])
        db.session.commit()
        mentee.mentor_id = mentor.id
        second_mentee.mentor_id = mentor.id
        third_mentee.mentor_id = mentor.id
        db.session.commit()
        mentor_id, mentee_id, second_mentee_id, third_mentee_id, other_mentee_id = (
            mentor.id,
            mentee.id,
            second_mentee.id,
            third_mentee.id,
            other_mentee.id,
        )

    client.post(
        "/login",
        data={"username": "schedule_mentor", "password": "ScheduleTest!2026", "role": "mentor"},
    )
    page = client.get("/mentor/meetings")
    assert page.status_code == 200
    assert b"Create a meeting" in page.data
    assert b"Taylor Mentee" in page.data
    assert b"All my active mentees" in page.data
    assert b"Jordan Mentee" in page.data
    assert b"Casey Mentee" in page.data
    assert b"Other Mentee" not in page.data

    invalid_choice = client.post(
        "/mentor/meetings",
        data={
            "audience": "selected",
            "mentee_ids": [str(other_mentee_id)],
            "date": str(date.today() + timedelta(days=2)),
            "time": "10:30 AM",
            "duration": "60",
            "mode": "online",
            "agenda": "Review the learning plan",
            "submit": "Create meeting",
        },
    )
    assert invalid_choice.status_code == 200
    with app.app_context():
        assert Meeting.query.count() == 0

    created = client.post(
        "/mentor/meetings",
        data={
            "audience": "selected",
            "mentee_ids": [str(mentee_id)],
            "date": str(date.today() + timedelta(days=2)),
            "time": "10:30 AM",
            "duration": "60",
            "mode": "online",
            "agenda": "Review the learning plan",
            "submit": "Create meeting",
        },
        follow_redirects=True,
    )
    assert created.status_code == 200
    assert b"Meeting created and added to the mentee&#39;s sessions." in created.data
    with app.app_context():
        meeting = Meeting.query.one()
        assert meeting.mentor_id == mentor_id
        assert meeting.mentee_id == mentee_id
        assert meeting.requested_by_id == mentor_id
        assert meeting.status == "Approved"
        assert Notification.query.filter_by(user_id=mentee_id, related_id=meeting.id).count() == 1

    selected_meetings = client.post(
        "/mentor/meetings",
        data={
            "audience": "selected",
            "mentee_ids": [str(mentee_id), str(second_mentee_id)],
            "date": str(date.today() + timedelta(days=3)),
            "time": "11:00 AM",
            "duration": "45",
            "mode": "online",
            "agenda": "Selected group progress check-in",
            "submit": "Create meeting",
        },
        follow_redirects=True,
    )
    assert selected_meetings.status_code == 200
    assert b"Meeting created for 2 selected mentees." in selected_meetings.data
    with app.app_context():
        selected_records = Meeting.query.filter_by(agenda="Selected group progress check-in").all()
        assert {meeting.mentee_id for meeting in selected_records} == {mentee_id, second_mentee_id}
        assert all(
            Notification.query.filter_by(user_id=meeting.mentee_id, related_id=meeting.id).count() == 1
            for meeting in selected_records
        )

    all_meetings = client.post(
        "/mentor/meetings",
        data={
            "audience": "all",
            "date": str(date.today() + timedelta(days=3)),
            "time": "2:00 PM",
            "duration": "45",
            "mode": "online",
            "agenda": "Group progress check-in",
            "submit": "Create meeting",
        },
        follow_redirects=True,
    )
    assert all_meetings.status_code == 200
    assert b"Meeting created for all 3 active mentees." in all_meetings.data
    with app.app_context():
        meetings = Meeting.query.order_by(Meeting.id).all()
        assert len(meetings) == 6
        group_sessions = Meeting.query.filter_by(agenda="Group progress check-in").all()
        group_session_mentee_ids = {meeting.mentee_id for meeting in group_sessions}
        assert group_session_mentee_ids == {mentee_id, second_mentee_id, third_mentee_id}
        for meeting in group_sessions:
            assert meeting.mentor_id == mentor_id
            assert meeting.status == "Approved"
            assert Notification.query.filter_by(user_id=meeting.mentee_id, related_id=meeting.id).count() == 1
        db.session.remove()
        db.engine.dispose()


def test_mentor_action_assignments_are_restricted_to_own_active_mentees(tmp_path):
    app = build_app(tmp_path)
    client = app.test_client()
    with app.app_context():
        mentor = User(username="task_owner", full_name="Task Owner", role="mentor", is_active=True)
        other_mentor = User(username="task_other", full_name="Other Mentor", role="mentor", is_active=True)
        own_mentee = User(username="task_own_student", full_name="Own Student", role="mentee", is_active=True)
        other_mentee = User(username="task_other_student", full_name="Other Student", role="mentee", is_active=True)
        inactive_mentee = User(username="task_inactive", full_name="Inactive Student", role="mentee", is_active=False)
        admin = User.query.filter_by(username="admin").one()
        for user in (mentor, other_mentor, own_mentee, other_mentee, inactive_mentee):
            user.set_password("ActionTest!2026")
        db.session.add_all([mentor, other_mentor, own_mentee, other_mentee, inactive_mentee])
        db.session.commit()
        own_mentee.mentor_id = mentor.id
        other_mentee.mentor_id = other_mentor.id
        inactive_mentee.mentor_id = mentor.id
        db.session.commit()
        mentor_id = mentor.id
        own_mentee_id = own_mentee.id
        other_mentee_id = other_mentee.id
        inactive_mentee_id = inactive_mentee.id
        admin_id = admin.id

    client.post(
        "/login",
        data={"username": "task_owner", "password": "ActionTest!2026", "role": "mentor"},
    )
    action_data = {
        "title": "Progress check",
        "description": "Submit a short progress update for the mentor.",
        "deadline": str(date.today() + timedelta(days=7)),
        "submit": "Create action",
    }
    for invalid_recipient in (str(other_mentee_id), str(inactive_mentee_id), str(admin_id), "not-a-number"):
        response = client.post(
            "/mentor/actions/new",
            data={**action_data, "mentee_ids": [invalid_recipient]},
        )
        assert response.status_code == 200
        error_page = response.data.lower()
        assert b"select from your mentee list" in error_page or b"only to your own active mentees" in error_page
        with app.app_context():
            assert Action.query.filter_by(title="Progress check").count() == 0
            assert ActionAssignment.query.count() == 0
            assert Notification.query.filter_by(type="action_assigned").count() == 0

    created = client.post(
        "/mentor/actions/new",
        data={**action_data, "mentee_ids": [str(own_mentee_id)]},
        follow_redirects=True,
    )
    assert created.status_code == 200
    with app.app_context():
        action = Action.query.filter_by(title="Progress check").one()
        assignment = ActionAssignment.query.one()
        assert action.created_by_id == mentor_id
        assert assignment.mentee_id == own_mentee_id
        assert Notification.query.filter_by(user_id=own_mentee_id, related_id=action.id).count() == 1
        assert Notification.query.filter_by(user_id=other_mentee_id, related_id=action.id).count() == 0
        db.session.remove()
        db.engine.dispose()


def test_notifications_are_capped_but_keep_the_unread_count(tmp_path):
    app = build_app(tmp_path)
    client = app.test_client()
    with app.app_context():
        user = User(username="notification_mentee", full_name="Nora Notifications", role="mentee")
        user.set_password("NotificationTest!2026")
        db.session.add(user)
        db.session.commit()
        db.session.add_all([
            Notification(user_id=user.id, type="update", message=f"Update {index}")
            for index in range(60)
        ])
        db.session.commit()

    client.post(
        "/login",
        data={"username": "notification_mentee", "password": "NotificationTest!2026", "role": "mentee"},
    )
    response = client.get("/api/notifications")
    assert response.status_code == 200
    assert len(response.json) == 50
    assert response.headers["X-Unread-Count"] == "60"
    for _ in range(51):
        response = client.get("/api/notifications")
        assert response.status_code == 200

    with app.app_context():
        db.session.remove()
        db.engine.dispose()


def test_mentor_can_manage_only_own_substitutions(tmp_path):
    app = build_app(tmp_path)
    client = app.test_client()
    with app.app_context():
        mentor = User(username="coverage_mentor", full_name="Morgan Primary", role="mentor")
        substitute = User(username="coverage_sub", full_name="Riley Substitute", role="mentor")
        other_mentor = User(username="coverage_other", full_name="Other Primary", role="mentor")
        mentee = User(username="coverage_student", full_name="Taylor Student", role="mentee")
        other_mentee = User(username="coverage_other_student", full_name="Other Student", role="mentee")
        for user in (mentor, substitute, other_mentor, mentee, other_mentee):
            user.set_password("CoverageTest!2026")
        db.session.add_all([mentor, substitute, other_mentor, mentee, other_mentee])
        db.session.commit()
        mentee.mentor_id = mentor.id
        other_mentee.mentor_id = other_mentor.id
        other_substitution = Substitution(
            mentee_id=other_mentee.id,
            mentor_id=other_mentor.id,
            substitute_mentor_id=substitute.id,
            reason="Other mentor coverage",
            start_date=date.today(),
            end_date=date.today() + timedelta(days=5),
            active=True,
        )
        db.session.add(other_substitution)
        db.session.commit()
        mentee_id = mentee.id
        substitute_id = substitute.id
        other_substitution_id = other_substitution.id

    client.post(
        "/login",
        data={"username": "coverage_mentor", "password": "CoverageTest!2026", "role": "mentor"},
    )
    page = client.get("/mentor/substitutions")
    assert page.status_code == 200
    assert b"every active mentee assigned to you" in page.data
    assert b"Other Student" not in page.data

    created = client.post(
        "/mentor/substitutions",
        data={
            "mentee_id": str(mentee_id),
            "substitute_mentor_id": str(substitute_id),
            "reason": "Primary mentor leave",
            "start_date": str(date.today() + timedelta(days=7)),
            "end_date": str(date.today() + timedelta(days=14)),
        },
        follow_redirects=True,
    )
    assert created.status_code == 200
    assert b"Substitution saved for all 1 active mentees." in created.data
    with app.app_context():
        substitution = Substitution.query.filter_by(mentee_id=mentee_id).one()
        substitution_id = substitution.id

    updated = client.post(
        f"/mentor/substitutions/{substitution_id}",
        data={
            "mentee_id": str(mentee_id),
            "substitute_mentor_id": str(substitute_id),
            "reason": "Updated coverage dates",
            "start_date": str(date.today() + timedelta(days=8)),
            "end_date": str(date.today() + timedelta(days=16)),
            "active": "y",
        },
        follow_redirects=True,
    )
    assert updated.status_code == 200
    assert b"Substitution updated." in updated.data
    with app.app_context():
        substitution = db.session.get(Substitution, substitution_id)
        assert substitution.reason == "Updated coverage dates"
        assert substitution.active is True

    forbidden = client.post(f"/mentor/substitutions/{other_substitution_id}", data={})
    assert forbidden.status_code == 404
    with app.app_context():
        other_substitution = db.session.get(Substitution, other_substitution_id)
        assert other_substitution.reason == "Other mentor coverage"
        db.session.remove()
        db.engine.dispose()


def test_mentor_inbox_is_limited_to_assigned_and_active_substitute_mentees(tmp_path):
    app = build_app(tmp_path)
    client = app.test_client()
    with app.app_context():
        mentor = User(username="scoped_mentor", full_name="Morgan Mentor", role="mentor", is_active=True)
        substitute = User(username="scoped_sub", full_name="Alex Substitute", role="mentor", is_active=True)
        other_mentor = User(username="scoped_other", full_name="Other Mentor", role="mentor", is_active=True)
        assigned = User(username="scoped_assigned", full_name="Assigned Mentee", role="mentee", is_active=True)
        covered = User(username="scoped_covered", full_name="Covered Mentee", role="mentee", is_active=True)
        unrelated = User(username="scoped_unrelated", full_name="Unrelated Mentee", role="mentee", is_active=True)
        for user in (mentor, substitute, other_mentor, assigned, covered, unrelated):
            user.set_password("ScopedTest!2026")
        db.session.add_all([mentor, substitute, other_mentor, assigned, covered, unrelated])
        db.session.commit()
        assigned.mentor_id = mentor.id
        covered.mentor_id = other_mentor.id
        unrelated.mentor_id = other_mentor.id
        db.session.add(Substitution(
            mentee_id=covered.id,
            mentor_id=other_mentor.id,
            substitute_mentor_id=mentor.id,
            reason="Temporary coverage",
            start_date=date.today(),
            end_date=date.today() + timedelta(days=2),
            active=True,
        ))
        db.session.commit()
        assigned_id, covered_id, unrelated_id = assigned.id, covered.id, unrelated.id

    client.post(
        "/login",
        data={"username": "scoped_mentor", "password": "ScopedTest!2026", "role": "mentor"},
    )
    inbox = client.get("/messages/")
    assert b"Assigned Mentee (mentee)" in inbox.data
    assert b"Covered Mentee (mentee)" in inbox.data
    assert b"Unrelated Mentee" not in inbox.data
    assert client.get(f"/messages/api/thread/{unrelated_id}").status_code == 404
    forbidden_send = client.post(
        "/messages/",
        data={"recipient_id": str(unrelated_id), "content": "Should not be sent"},
        follow_redirects=True,
    )
    assert b"This recipient is not available for your account." in forbidden_send.data
    with app.app_context():
        assert Message.query.count() == 0
        assert assigned_id != covered_id
        db.session.remove()
        db.engine.dispose()


def test_mentee_can_update_only_own_task_progress(tmp_path):
    app = build_app(tmp_path)
    client = app.test_client()
    with app.app_context():
        mentor = User(username="progress_mentor", full_name="Progress Mentor", role="mentor", is_active=True)
        mentee = User(username="progress_mentee", full_name="Progress Mentee", role="mentee", is_active=True)
        other_mentee = User(username="progress_other", full_name="Other Mentee", role="mentee", is_active=True)
        for user in (mentor, mentee, other_mentee):
            user.set_password("ProgressTest!2026")
        db.session.add_all([mentor, mentee, other_mentee])
        db.session.commit()
        action = Action(
            title="Weekly reflection",
            description="Submit this week's reflection.",
            deadline=date.today() + timedelta(days=5),
            created_by_id=mentor.id,
        )
        db.session.add(action)
        db.session.flush()
        own_assignment = ActionAssignment(action_id=action.id, mentee_id=mentee.id)
        other_assignment = ActionAssignment(action_id=action.id, mentee_id=other_mentee.id)
        completed_action = Action(
            title="Completed orientation",
            description="Finish the orientation activity.",
            deadline=date.today() - timedelta(days=1),
            created_by_id=mentor.id,
        )
        db.session.add(completed_action)
        db.session.flush()
        completed_assignment = ActionAssignment(
            action_id=completed_action.id,
            mentee_id=mentee.id,
            status="Completed",
            progress=0,
        )
        db.session.add_all([own_assignment, other_assignment, completed_assignment])
        db.session.commit()
        own_assignment_id, other_assignment_id = own_assignment.id, other_assignment.id

    client.post(
        "/login",
        data={"username": "progress_mentee", "password": "ProgressTest!2026", "role": "mentee"},
    )
    response = client.post(
        f"/mentee/actions/{own_assignment_id}",
        data={"status": "In Progress", "progress": "45", "notes": "Draft is underway."},
        follow_redirects=True,
    )
    assert response.status_code == 200
    dashboard = client.get("/mentee/dashboard")
    assert b"Student progress report" in dashboard.data
    assert b"Overall task progress" in dashboard.data
    assert b"73%" in dashboard.data
    assert b"Total tasks" in dashboard.data
    assert b"Completed orientation" in dashboard.data
    assert b"Draft is underway." in dashboard.data
    with app.app_context():
        assignment = db.session.get(ActionAssignment, own_assignment_id)
        assert assignment.status == "In Progress"
        assert assignment.progress == 45
        assert assignment.notes == "Draft is underway."

    assert client.post(
        f"/mentee/actions/{own_assignment_id}",
        data={"status": "Completed", "progress": "150", "notes": ""},
    ).status_code == 302
    assert client.post(
        f"/mentee/actions/{other_assignment_id}",
        data={"status": "Completed", "progress": "100", "notes": ""},
    ).status_code == 404
    with app.app_context():
        assignment = db.session.get(ActionAssignment, own_assignment_id)
        assert assignment.status == "In Progress"
        assert assignment.progress == 45
        db.session.remove()
        db.engine.dispose()


def test_admin_allocation_validates_roles_and_audits_actual_admin(tmp_path):
    app = build_app(tmp_path)
    client = app.test_client()
    with app.app_context():
        default_admin = User.query.filter_by(username="admin").one()
        default_admin.must_change_password = False
        second_admin = User(username="allocation_admin", full_name="Allocation Admin", role="admin", is_active=True)
        mentor = User(username="valid_mentor", full_name="Valid Mentor", role="mentor", is_active=True)
        mentee = User(username="valid_mentee", full_name="Valid Mentee", role="mentee", is_active=True)
        for user in (second_admin, mentor, mentee):
            user.set_password("AllocateTest!2026")
        db.session.add_all([second_admin, mentor, mentee])
        db.session.commit()
        admin_id, mentor_id, mentee_id, invalid_role_id = (
            second_admin.id,
            mentor.id,
            mentee.id,
            default_admin.id,
        )

    client.post(
        "/login",
        data={"username": "allocation_admin", "password": "AllocateTest!2026", "role": "admin"},
    )
    for raw_id in ("garbage", str(invalid_role_id)):
        assert client.post(f"/admin/allocations/{mentee_id}", data={"mentor_id": raw_id}).status_code == 302
        with app.app_context():
            assert AllocationHistory.query.count() == 0

    client.post(f"/admin/allocations/{mentee_id}", data={"mentor_id": str(mentor_id)})
    with app.app_context():
        history = AllocationHistory.query.one()
        assert history.changed_by_id == admin_id
        assert history.mentee_id == mentee_id
        assert AuditLog.query.filter_by(user_id=admin_id, action="mentor_assignment").count() == 1
        db.session.remove()
        db.engine.dispose()


def test_admin_substitution_rejects_unassigned_past_and_overlapping_dates(tmp_path):
    app = build_app(tmp_path)
    client = app.test_client()
    with app.app_context():
        admin = User.query.filter_by(username="admin").one()
        admin.must_change_password = False
        mentor = User(username="sub_owner", full_name="Current Mentor", role="mentor", is_active=True)
        substitute = User(username="substitute_mentor", full_name="Cover Mentor", role="mentor", is_active=True)
        mentee = User(username="sub_mentee", full_name="Assigned Mentee", role="mentee", is_active=True)
        unassigned = User(username="sub_unassigned", full_name="Unassigned Mentee", role="mentee", is_active=True)
        for user in (mentor, substitute, mentee, unassigned):
            user.set_password("SubTest!2026")
        db.session.add_all([mentor, substitute, mentee, unassigned])
        db.session.commit()
        mentee.mentor_id = mentor.id
        db.session.add(Substitution(
            mentee_id=mentee.id,
            mentor_id=mentor.id,
            substitute_mentor_id=substitute.id,
            reason="Existing coverage",
            start_date=date.today() + timedelta(days=3),
            end_date=date.today() + timedelta(days=6),
            active=True,
        ))
        db.session.commit()
        mentee_id, unassigned_id, substitute_id = mentee.id, unassigned.id, substitute.id
        admin_password = app.config["SEED_ADMIN_PASSWORD"]

    client.post("/login", data={"username": "admin", "password": admin_password, "role": "admin"})
    base_data = {
        "substitute_mentor_id": str(substitute_id),
        "reason": "Temporary cover",
        "start_date": str(date.today() + timedelta(days=1)),
        "end_date": str(date.today() + timedelta(days=2)),
    }
    invalid_cases = [
        {**base_data, "mentee_id": str(unassigned_id)},
        {**base_data, "mentee_id": str(mentee_id), "start_date": str(date.today() - timedelta(days=1))},
        {
            **base_data,
            "mentee_id": str(mentee_id),
            "start_date": str(date.today() + timedelta(days=4)),
            "end_date": str(date.today() + timedelta(days=5)),
        },
    ]
    for data in invalid_cases:
        response = client.post("/admin/substitutions", data=data, follow_redirects=True)
        assert response.status_code == 200
    with app.app_context():
        assert Substitution.query.count() == 1
        db.session.remove()
        db.engine.dispose()


def test_admin_cannot_deactivate_self_or_last_active_admin(tmp_path):
    app = build_app(tmp_path)
    client = app.test_client()
    with app.app_context():
        admin = User.query.filter_by(username="admin").one()
        admin.must_change_password = False
        admin_id = admin.id
        password = app.config["SEED_ADMIN_PASSWORD"]
        db.session.commit()
    client.post("/login", data={"username": "admin", "password": password, "role": "admin"})
    response = client.post(f"/admin/users/{admin_id}/toggle", follow_redirects=True)
    assert response.status_code == 200
    assert b"You cannot deactivate the administrator account" in response.data
    with app.app_context():
        assert db.session.get(User, admin_id).is_active is True
        db.session.remove()
        db.engine.dispose()


def test_mentee_meeting_requests_reject_past_dates_and_invalid_durations(tmp_path):
    app = build_app(tmp_path)
    client = app.test_client()
    with app.app_context():
        mentor = User(username="date_mentor", full_name="Date Mentor", role="mentor", is_active=True)
        mentee = User(username="date_mentee", full_name="Date Mentee", role="mentee", is_active=True)
        mentor.set_password("MeetingTest!2026")
        mentee.set_password("MeetingTest!2026")
        db.session.add_all([mentor, mentee])
        db.session.commit()
        mentee.mentor_id = mentor.id
        db.session.commit()

    client.post(
        "/login",
        data={"username": "date_mentee", "password": "MeetingTest!2026", "role": "mentee"},
    )
    form_data = {
        "date": str(date.today() - timedelta(days=1)),
        "time": "10:30 AM",
        "duration": "60",
        "mode": "online",
        "agenda": "Review learning goals",
    }
    for data in (
        form_data,
        {**form_data, "date": str(date.today() + timedelta(days=1)), "duration": "5"},
        {**form_data, "date": str(date.today() + timedelta(days=1)), "duration": "300"},
    ):
        assert client.post("/mentee/meetings/new", data=data).status_code == 200
    with app.app_context():
        assert Meeting.query.count() == 0
        db.session.remove()
        db.engine.dispose()


def test_concern_threads_allow_scoped_replies_and_status_updates(tmp_path):
    app = build_app(tmp_path)
    mentee_client = app.test_client()
    mentor_client = app.test_client()
    with app.app_context():
        admin = User.query.filter_by(username="admin").one()
        admin.must_change_password = False
        mentor = User(username="concern_mentor", full_name="Concern Mentor", role="mentor", is_active=True)
        other_mentor = User(username="concern_other", full_name="Other Mentor", role="mentor", is_active=True)
        mentee = User(username="concern_mentee", full_name="Concern Mentee", role="mentee", is_active=True)
        for user in (mentor, other_mentor, mentee):
            user.set_password("ConcernTest!2026")
        db.session.add_all([mentor, other_mentor, mentee])
        db.session.commit()
        mentee.mentor_id = mentor.id
        concern = Concern(
            mentee_id=mentee.id,
            mentor_id=mentor.id,
            title="Course support",
            category="academic",
            description="I need support with the project.",
            priority="medium",
            status="Open",
        )
        db.session.add(concern)
        db.session.commit()
        concern_id, admin_password = concern.id, app.config["SEED_ADMIN_PASSWORD"]

    mentee_client.post(
        "/login",
        data={"username": "concern_mentee", "password": "ConcernTest!2026", "role": "mentee"},
    )
    mentor_client.post(
        "/login",
        data={"username": "concern_mentor", "password": "ConcernTest!2026", "role": "mentor"},
    )
    assert mentee_client.post(
        f"/mentee/concerns/{concern_id}/reply",
        data={"response": "Thank you for following up."},
    ).status_code == 302
    assert mentor_client.post(
        f"/mentor/concerns/{concern_id}",
        data={"status": "In Progress", "response": "Let's set a plan."},
    ).status_code == 302
    assert mentor_client.post(
        f"/mentor/concerns/{concern_id}",
        data={"status": "Resolved", "response": ""},
    ).status_code == 302
    with app.app_context():
        updated = db.session.get(Concern, concern_id)
        assert updated.status == "Resolved"
        assert "Thank you for following up." in updated.comments
        assert "Let's set a plan." in updated.comments
        assert Notification.query.filter_by(user_id=updated.mentee_id, type="concern_update").count() == 2
        db.session.remove()
        db.engine.dispose()


def test_mentor_can_create_substitution_for_all_assigned_mentees_atomically(tmp_path):
    app = build_app(tmp_path)
    client = app.test_client()
    with app.app_context():
        mentor = User(username="bulk_coverage_mentor", full_name="Morgan Mentor", role="mentor", is_active=True)
        substitute = User(username="bulk_coverage_sub", full_name="Riley Substitute", role="mentor", is_active=True)
        mentees = [
            User(username=f"bulk_student_{index}", full_name=f"Bulk Student {index}", role="mentee", is_active=True)
            for index in range(3)
        ]
        for user in [mentor, substitute, *mentees]:
            user.set_password("BulkCoverage!2026")
        db.session.add_all([mentor, substitute, *mentees])
        db.session.commit()
        for mentee in mentees:
            mentee.mentor_id = mentor.id
        db.session.add(Substitution(
            mentee_id=mentees[0].id,
            mentor_id=mentor.id,
            substitute_mentor_id=substitute.id,
            reason="Existing overlap",
            start_date=date.today() + timedelta(days=2),
            end_date=date.today() + timedelta(days=4),
            active=True,
        ))
        db.session.commit()
        mentee_ids = {mentee.id for mentee in mentees}
        substitute_id = substitute.id

    client.post(
        "/login",
        data={
            "username": "bulk_coverage_mentor",
            "password": "BulkCoverage!2026",
            "role": "mentor",
        },
    )
    page = client.get("/mentor/substitutions")
    assert b"every active mentee assigned to you" in page.data
    assert b"Bulk Student 0" in page.data
    overlap = client.post(
        "/mentor/substitutions",
        data={
            "substitute_mentor_id": str(substitute_id),
            "reason": "Shared cover",
            "start_date": str(date.today() + timedelta(days=3)),
            "end_date": str(date.today() + timedelta(days=5)),
        },
        follow_redirects=True,
    )
    assert overlap.status_code == 200
    assert b"No substitutions were created because active coverage overlaps" in overlap.data
    with app.app_context():
        assert Substitution.query.count() == 1

    created = client.post(
        "/mentor/substitutions",
        data={
            "substitute_mentor_id": str(substitute_id),
            "reason": "Shared cover for mentor leave",
            "start_date": str(date.today() + timedelta(days=7)),
            "end_date": str(date.today() + timedelta(days=12)),
        },
        follow_redirects=True,
    )
    assert created.status_code == 200
    assert b"Substitution saved for all 3 active mentees." in created.data
    with app.app_context():
        created_records = Substitution.query.filter(
            Substitution.start_date == date.today() + timedelta(days=7)
        ).all()
        assert {record.mentee_id for record in created_records} == mentee_ids
        assert all(record.substitute_mentor_id == substitute_id for record in created_records)
        assert Notification.query.filter_by(type="substitution").count() == 3
        audit = AuditLog.query.filter_by(action="mentor_substitution_create").one()
        assert all(mentee.username in audit.details for mentee in User.query.filter(User.id.in_(mentee_ids)).all())
        db.session.remove()
        db.engine.dispose()


def _student_workbook(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["username", "full_name", "email", "mentor_username"])
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def test_admin_imports_students_from_excel_and_downloads_credentials(tmp_path):
    app = build_app(tmp_path)
    client = app.test_client()
    with app.app_context():
        admin = User.query.filter_by(username="admin").one()
        admin.must_change_password = False
        mentor = User(username="excel_mentor", full_name="Alex Mentor", role="mentor", is_active=True)
        mentor.set_password("ExcelMentorTest!2026")
        db.session.add(mentor)
        db.session.commit()
        admin_password = app.config["SEED_ADMIN_PASSWORD"]
        mentor_id = mentor.id

    client.post(
        "/login",
        data={"username": "admin", "password": admin_password, "role": "admin"},
    )
    template_response = client.get("/admin/users/import/template.xlsx")
    assert template_response.status_code == 200
    template = load_workbook(BytesIO(template_response.data), read_only=True)
    assert next(template.active.iter_rows(values_only=True)) == (
        "username",
        "full_name",
        "email",
        "mentor_username",
    )
    template.close()

    response = client.post(
        "/admin/users/import",
        data={
            "file": (
                _student_workbook([[
                    "excel_student",
                    "Jamie Example",
                    "jamie@example.com",
                    "excel_mentor",
                ]]),
                "students.xlsx",
            )
        },
        content_type="multipart/form-data",
    )
    assert response.status_code == 200
    assert response.mimetype == "text/csv"
    credentials = next(csv.DictReader(StringIO(response.data.decode("utf-8-sig"))))
    assert credentials["username"] == "excel_student"
    assert credentials["temporary_password"]
    temporary_password = credentials["temporary_password"]
    assert temporary_password[0] not in ("=", "+", "-", "@")
    assert any(character.isupper() for character in temporary_password)
    assert any(character.islower() for character in temporary_password)
    assert any(character.isdigit() for character in temporary_password)
    assert any(not character.isalnum() for character in temporary_password)
    with app.app_context():
        student = User.query.filter_by(username="excel_student").one()
        assert student.role == "mentee"
        assert student.mentor_id == mentor_id
        assert student.must_change_password is True
        assert student.check_password(temporary_password)
        assert student.password_hash != temporary_password
        db.session.remove()
        db.engine.dispose()


def test_admin_student_import_rejects_invalid_rows_without_partial_creation(tmp_path):
    app = build_app(tmp_path)
    client = app.test_client()
    with app.app_context():
        admin = User.query.filter_by(username="admin").one()
        admin.must_change_password = False
        db.session.commit()
        admin_password = app.config["SEED_ADMIN_PASSWORD"]

    client.post(
        "/login",
        data={"username": "admin", "password": admin_password, "role": "admin"},
    )
    response = client.post(
        "/admin/users/import",
        data={
            "file": (
                _student_workbook([
                    ["would_be_created", "Valid Student", "valid@example.test", ""],
                    ["bad_student", "Invalid Student", "not-an-email", "missing_mentor"],
                ]),
                "students.xlsx",
            )
        },
        content_type="multipart/form-data",
    )
    assert response.status_code == 200
    assert b"No accounts were imported." in response.data
    assert b"Row 3" in response.data
    with app.app_context():
        assert User.query.filter_by(username="would_be_created").first() is None
        assert User.query.filter_by(username="bad_student").first() is None
        db.session.remove()
        db.engine.dispose()
