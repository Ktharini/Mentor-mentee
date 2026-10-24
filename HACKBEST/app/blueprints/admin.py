import csv
import re
from datetime import date, datetime
from io import BytesIO, StringIO
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError

from email_validator import EmailNotValidError, validate_email
from flask import Blueprint, flash, make_response, redirect, render_template, request, send_file, url_for
from flask_login import current_user
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import func
from sqlalchemy.exc import IntegrityError

from app.auth_security import invalidate_user_sessions
from app.forms import AllocationForm, ConcernReviewForm, SubstitutionForm, UserForm
from app.models import (
    Action,
    ActionAssignment,
    AllocationHistory,
    AuditLog,
    Concern,
    Meeting,
    Message,
    MessageDelivery,
    Notification,
    PasswordResetToken,
    Substitution,
    User,
    UserAuthState,
    db,
)
from app.security import role_required
from app.utils import add_notification, append_concern_response, generate_strong_password

admin_bp = Blueprint("admin", __name__)
STUDENT_IMPORT_MAX_BYTES = 5 * 1024 * 1024
STUDENT_IMPORT_MAX_ROWS = 500
STUDENT_USERNAME_PATTERN = re.compile(r"^[A-Za-z0-9._-]{3,50}$")


def _temporary_student_password():
    while True:
        password = generate_strong_password(16)
        if (
            password[0] not in ("=", "+", "-", "@")
            and any(character.isupper() for character in password)
            and any(character.islower() for character in password)
            and any(character.isdigit() for character in password)
            and any(not character.isalnum() for character in password)
        ):
            return password


def _safe_csv_cell(value):
    text = str(value)
    return f"'{text}" if text.startswith(("=", "+", "-", "@", "\t", "\r")) else text


def _users_response(temporary_password=None):
    response = make_response(render_template(
        "admin/users.html",
        users=User.query.order_by(User.created_at.desc()).all(),
        temporary_password=temporary_password,
        active_admin_count=User.query.filter_by(role="admin", is_active=True).count(),
    ))
    if temporary_password:
        response.headers["Cache-Control"] = "no-store, max-age=0"
        response.headers["Pragma"] = "no-cache"
    return response


@admin_bp.route("/")
@role_required("admin")
def dashboard():
    mentors = User.query.filter_by(role="mentor", is_active=True).all()
    mentees = User.query.filter_by(role="mentee", is_active=True).all()
    recent_logs = AuditLog.query.order_by(AuditLog.created_at.desc()).limit(10).all()
    today = date.today()
    current_month = today.replace(day=1)
    sessions_this_month = Meeting.query.filter(
        Meeting.date >= current_month,
        Meeting.date <= today,
        Meeting.status.notin_(["Cancelled", "Rejected"]),
    ).count()
    pending_approvals = Meeting.query.filter_by(status="Pending").count()

    month_starts = []
    for offset in range(5, -1, -1):
        month_index = today.year * 12 + today.month - 1 - offset
        month_starts.append(date(month_index // 12, month_index % 12 + 1, 1))
    growth_labels = [month.strftime("%b") for month in month_starts]
    user_growth = []
    meeting_activity = []
    for index, month_start in enumerate(month_starts):
        next_month = month_starts[index + 1] if index + 1 < len(month_starts) else date(
            month_start.year + (month_start.month == 12),
            1 if month_start.month == 12 else month_start.month + 1,
            1,
        )
        month_datetime = datetime.combine(month_start, datetime.min.time())
        next_month_datetime = datetime.combine(next_month, datetime.min.time())
        user_growth.append(
            User.query.filter(User.created_at >= month_datetime, User.created_at < next_month_datetime).count()
        )
        meeting_activity.append(
            Meeting.query.filter(Meeting.date >= month_start, Meeting.date < next_month).count()
        )

    from app.models import Concern

    pending_concerns = Concern.query.filter(Concern.status.notin_(["Resolved", "Closed"])).count()
    return render_template(
        "admin/dashboard.html",
        mentors=mentors,
        mentees=mentees,
        recent_logs=recent_logs,
        total_users=User.query.count(),
        sessions_this_month=sessions_this_month,
        pending_approvals=pending_approvals,
        pending_concerns=pending_concerns,
        growth_labels=growth_labels,
        user_growth=user_growth,
        meeting_activity=meeting_activity,
    )


@admin_bp.route("/users")
@role_required("admin")
def users():
    return _users_response()


@admin_bp.route("/users/import", methods=["GET", "POST"])
@role_required("admin")
def import_students():
    if request.method == "GET":
        return render_template("admin/student_import.html", import_errors=[])

    workbook_file = request.files.get("file")
    if not workbook_file or not workbook_file.filename:
        return render_template(
            "admin/student_import.html",
            import_errors=["Choose an .xlsx workbook to upload."],
        )
    if not workbook_file.filename.lower().endswith(".xlsx"):
        return render_template(
            "admin/student_import.html",
            import_errors=["Only .xlsx workbooks are supported."],
        )

    workbook_bytes = workbook_file.read(STUDENT_IMPORT_MAX_BYTES + 1)
    if len(workbook_bytes) > STUDENT_IMPORT_MAX_BYTES:
        return render_template(
            "admin/student_import.html",
            import_errors=["The workbook must be 5 MB or smaller."],
        )

    try:
        workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=True)
        sheet = workbook.active
        if sheet.max_row is None or sheet.max_row < 1:
            workbook.close()
            return render_template(
                "admin/student_import.html",
                import_errors=["The workbook is empty."],
            )
        if sheet.max_row > STUDENT_IMPORT_MAX_ROWS + 1:
            workbook.close()
            return render_template(
                "admin/student_import.html",
                import_errors=[f"A workbook can contain at most {STUDENT_IMPORT_MAX_ROWS} student rows."],
            )
        workbook_rows = list(sheet.iter_rows(values_only=True))
        workbook.close()
    except (BadZipFile, InvalidFileException, ParseError, OSError, ValueError, KeyError, TypeError):
        return render_template(
            "admin/student_import.html",
            import_errors=["This file is not a valid .xlsx workbook."],
        )

    headers = [
        str(value).strip().lower().replace(" ", "_") if value is not None else ""
        for value in workbook_rows[0]
    ]
    allowed_headers = {"username", "full_name", "email", "mentor_username"}
    header_errors = []
    if any(header and header not in allowed_headers for header in headers):
        header_errors.append("Use only these columns: username, full_name, email, mentor_username.")
    if len([header for header in headers if header]) != len(set(header for header in headers if header)):
        header_errors.append("Column names must not be duplicated.")
    missing_headers = {"username", "full_name", "email"} - set(headers)
    if missing_headers:
        header_errors.append(f"Required columns are missing: {', '.join(sorted(missing_headers))}.")
    if header_errors:
        return render_template("admin/student_import.html", import_errors=header_errors)

    columns = {header: index for index, header in enumerate(headers) if header}
    existing_usernames, existing_emails = set(), set()
    for username, email in db.session.query(func.lower(User.username), func.lower(User.email)).all():
        existing_usernames.add(username)
        if email:
            existing_emails.add(email)
    available_mentors = User.query.filter_by(role="mentor", is_active=True).all()
    mentors_by_username = {mentor.username.casefold(): mentor for mentor in available_mentors}
    seen_usernames, seen_emails = set(), set()
    valid_rows, row_errors = [], []

    def cell_value(row, column):
        index = columns.get(column)
        if index is None or index >= len(row) or row[index] is None:
            return ""
        return str(row[index]).strip()

    for row_number, row in enumerate(workbook_rows[1:], start=2):
        if not any(value is not None and str(value).strip() for value in row):
            continue

        username = cell_value(row, "username")
        full_name = cell_value(row, "full_name")
        email_input = cell_value(row, "email")
        mentor_username = cell_value(row, "mentor_username")
        errors = []
        if not STUDENT_USERNAME_PATTERN.fullmatch(username):
            errors.append("username must be 3–50 letters, numbers, dots, underscores, or hyphens.")
        if len(full_name) < 2 or len(full_name) > 120 or any(ord(char) < 32 for char in full_name):
            errors.append("full_name must be 2–120 characters without control characters.")
        normalized_email = ""
        try:
            normalized_email = validate_email(email_input, check_deliverability=False).normalized.lower()
            if len(normalized_email) > 120:
                errors.append("email must be at most 120 characters.")
        except EmailNotValidError:
            errors.append("email must be a valid email address.")
        if username:
            username_key = username.casefold()
            if username_key in existing_usernames or username_key in seen_usernames:
                errors.append("username already exists.")
            seen_usernames.add(username_key)
        if normalized_email:
            email_key = normalized_email.casefold()
            if email_key in existing_emails or email_key in seen_emails:
                errors.append("email already exists.")
            seen_emails.add(email_key)

        mentor = None
        if mentor_username:
            mentor = mentors_by_username.get(mentor_username.casefold())
            if not mentor:
                errors.append("mentor_username must match an active mentor.")

        if errors:
            row_errors.append(f"Row {row_number}: {' '.join(errors)}")
        else:
            valid_rows.append({
                "username": username,
                "full_name": full_name,
                "email": normalized_email,
                "mentor": mentor,
            })

    if not valid_rows and not row_errors:
        row_errors.append("Add at least one student row beneath the header.")
    if row_errors:
        shown_errors = row_errors[:50]
        if len(row_errors) > len(shown_errors):
            shown_errors.append(f"{len(row_errors) - len(shown_errors)} additional row errors were omitted.")
        return render_template("admin/student_import.html", import_errors=shown_errors)

    credential_rows = []
    generated_passwords = set()
    for row in valid_rows:
        password = _temporary_student_password()
        while password in generated_passwords:
            password = _temporary_student_password()
        generated_passwords.add(password)
        user = User(
            username=row["username"],
            full_name=row["full_name"],
            email=row["email"],
            role="mentee",
            mentor_id=row["mentor"].id if row["mentor"] else None,
            is_active=True,
            must_change_password=True,
        )
        user.set_password(password)
        db.session.add(user)
        credential_rows.append([
            row["username"],
            row["full_name"],
            row["email"],
            row["mentor"].username if row["mentor"] else "",
            password,
        ])

    db.session.add(Notification(
        user_id=current_user.id,
        type="user_import",
        message=f"Imported {len(valid_rows)} mentee accounts from a spreadsheet.",
    ))
    db.session.add(AuditLog(
        user_id=current_user.id,
        action="bulk_mentee_import",
        details=f"Imported {len(valid_rows)} mentee accounts.",
    ))
    try:
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        return render_template(
            "admin/student_import.html",
            import_errors=["An account with one of these usernames or email addresses was just created. Refresh the workbook and try again."],
        )

    report = StringIO(newline="")
    csv_writer = csv.writer(report)
    csv_writer.writerow(["username", "full_name", "email", "mentor_username", "temporary_password"])
    for row in credential_rows:
        csv_writer.writerow([_safe_csv_cell(value) for value in row])
    response = send_file(
        BytesIO(report.getvalue().encode("utf-8-sig")),
        mimetype="text/csv",
        as_attachment=True,
        download_name="proguide-student-credentials.csv",
    )
    response.headers["Cache-Control"] = "no-store, max-age=0"
    response.headers["Pragma"] = "no-cache"
    return response


@admin_bp.route("/users/import/template.xlsx")
@role_required("admin")
def student_import_template():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Students"
    sheet.append(["username", "full_name", "email", "mentor_username"])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:D1"
    for column, width in {"A": 24, "B": 30, "C": 34, "D": 24}.items():
        sheet.column_dimensions[column].width = width
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="proguide-student-import-template.xlsx",
    )


@admin_bp.route("/users/new", methods=["GET", "POST"])
@role_required("admin")
def create_user():
    form = UserForm()
    mentors = User.query.filter_by(role="mentor", is_active=True).all()
    form.role.choices = [("mentor", "Mentor"), ("mentee", "Mentee")]
    if form.validate_on_submit():
        if User.query.filter_by(username=form.username.data.strip()).first():
            flash("Username already exists.", "danger")
            return render_template("admin/user_form.html", form=form, mentors=mentors)
        password = form.password.data or _temporary_student_password()
        user = User(
            username=form.username.data.strip(),
            full_name=form.full_name.data.strip(),
            email=(form.email.data or "").strip() or None,
            role=form.role.data,
            is_active=True,
            must_change_password=True,
        )
        user.set_password(password)
        db.session.add(user)
        db.session.commit()
        if user.role == "mentee":
            admin = User.query.filter_by(role="admin").first()
            if admin:
                Notification.add(admin.id, "user_created", f"New mentee account created for {user.full_name}.")
        AuditLog.log(user.id, "user_create", f"User {user.username} created by admin.")
        flash("User created. Share the temporary password securely; it will not be shown again.", "success")
        return _users_response(temporary_password=password)
    return render_template("admin/user_form.html", form=form, mentors=mentors)


@admin_bp.route("/users/<int:user_id>/toggle", methods=["POST"])
@role_required("admin")
def toggle_user(user_id):
    user = User.query.get_or_404(user_id)
    if user.id == current_user.id and user.is_active:
        flash("You cannot deactivate the administrator account you are currently using.", "danger")
        return redirect(url_for("admin.users"))
    if user.role == "admin" and user.is_active:
        active_admins = User.query.filter_by(role="admin", is_active=True).count()
        if active_admins <= 1:
            flash("You cannot deactivate the last active administrator.", "danger")
            return redirect(url_for("admin.users"))
    user.is_active = not user.is_active
    invalidate_user_sessions(user.id)
    db.session.commit()
    AuditLog.log(user.id, "user_toggle", f"Account status changed to {'active' if user.is_active else 'inactive'}.")
    flash("User status updated.", "success")
    return redirect(url_for("admin.users"))


@admin_bp.route("/users/<int:user_id>/delete", methods=["POST"])
@role_required("admin")
def delete_user(user_id):
    user = User.query.get_or_404(user_id)
    if user.id == current_user.id:
        flash("You cannot delete the account you are currently using.", "danger")
        return redirect(url_for("admin.users"))
    if user.role == "admin" and user.is_active:
        active_admins = User.query.filter_by(role="admin", is_active=True).count()
        if active_admins <= 1:
            flash("You cannot delete the last active administrator.", "danger")
            return redirect(url_for("admin.users"))

    deleted_name = user.full_name
    user_id = user.id
    owned_action_ids = [
        action_id
        for (action_id,) in db.session.query(Action.id).filter_by(created_by_id=user_id).all()
    ]
    related_message_ids = [
        message_id
        for (message_id,) in db.session.query(Message.id).filter(
            (Message.sender_id == user_id) | (Message.recipient_id == user_id)
        ).all()
    ]

    if related_message_ids:
        MessageDelivery.query.filter(MessageDelivery.message_id.in_(related_message_ids)).delete(
            synchronize_session=False
        )
        Message.query.filter(Message.id.in_(related_message_ids)).delete(synchronize_session=False)
    if owned_action_ids:
        ActionAssignment.query.filter(ActionAssignment.action_id.in_(owned_action_ids)).delete(
            synchronize_session=False
        )
    ActionAssignment.query.filter_by(mentee_id=user_id).delete(synchronize_session=False)
    Action.query.filter_by(created_by_id=user_id).delete(synchronize_session=False)
    Meeting.query.filter(
        (Meeting.mentee_id == user_id)
        | (Meeting.mentor_id == user_id)
        | (Meeting.requested_by_id == user_id)
    ).delete(synchronize_session=False)
    Concern.query.filter_by(mentee_id=user_id).delete(synchronize_session=False)
    Concern.query.filter_by(mentor_id=user_id).update({"mentor_id": None}, synchronize_session=False)
    Substitution.query.filter(
        (Substitution.mentee_id == user_id)
        | (Substitution.mentor_id == user_id)
        | (Substitution.substitute_mentor_id == user_id)
    ).delete(synchronize_session=False)
    AllocationHistory.query.filter(
        (AllocationHistory.mentee_id == user_id)
        | (AllocationHistory.old_mentor_id == user_id)
        | (AllocationHistory.new_mentor_id == user_id)
        | (AllocationHistory.changed_by_id == user_id)
    ).delete(synchronize_session=False)
    User.query.filter_by(mentor_id=user_id).update({"mentor_id": None}, synchronize_session=False)
    Notification.query.filter_by(user_id=user_id).delete(synchronize_session=False)
    PasswordResetToken.query.filter_by(user_id=user_id).delete(synchronize_session=False)
    UserAuthState.query.filter_by(user_id=user_id).delete(synchronize_session=False)
    AuditLog.query.filter_by(user_id=user_id).delete(synchronize_session=False)

    db.session.delete(user)
    db.session.add(
        AuditLog(
            user_id=current_user.id,
            action="user_delete",
            details=f"Deleted user {deleted_name} (ID {user_id}) and related records.",
        )
    )
    db.session.commit()
    flash(f"Deleted {deleted_name} and their related records.", "success")
    return redirect(url_for("admin.users"))


@admin_bp.route("/users/<int:user_id>/reset-password", methods=["POST"])
@role_required("admin")
def reset_password(user_id):
    user = User.query.get_or_404(user_id)
    new_password = _temporary_student_password()
    user.set_password(new_password)
    user.must_change_password = True
    invalidate_user_sessions(user.id)
    db.session.commit()
    Notification.add(user.id, "password_reset", "Your password has been reset. Please set a new one.")
    AuditLog.log(user.id, "password_reset", "Password reset by admin.")
    flash("Password reset. Share the temporary password securely; it will not be shown again.", "success")
    return _users_response(temporary_password=new_password)


@admin_bp.route("/allocations")
@role_required("admin")
def allocations():
    mentors = User.query.filter_by(role="mentor", is_active=True).order_by(User.full_name).all()
    mentees = User.query.filter_by(role="mentee", is_active=True).order_by(User.full_name).all()
    assigned_mentor_ids = {mentee.mentor_id for mentee in mentees if mentee.mentor_id is not None}
    current_mentors = {
        mentor.id: mentor.full_name
        for mentor in User.query.filter(
            User.id.in_(assigned_mentor_ids),
            User.role == "mentor",
        ).all()
    } if assigned_mentor_ids else {}
    return render_template(
        "admin/allocations.html",
        mentors=mentors,
        mentees=mentees,
        current_mentors=current_mentors,
    )


@admin_bp.route("/allocations/<int:mentee_id>", methods=["POST"])
@role_required("admin")
def assign_mentor(mentee_id):
    raw_mentor_id = request.form.get("mentor_id", "").strip()
    try:
        mentor_id = int(raw_mentor_id)
    except (TypeError, ValueError):
        flash("Missing mentor selection.", "danger")
        return redirect(url_for("admin.allocations"))
    mentee = User.query.filter_by(id=mentee_id, role="mentee", is_active=True).first()
    mentor = User.query.filter_by(id=mentor_id, role="mentor", is_active=True).first()
    if not mentee or not mentor:
        flash("Choose an active mentee and an active mentor.", "danger")
        return redirect(url_for("admin.allocations"))
    old_mentor = mentee.mentor_id
    mentee.mentor_id = mentor.id
    db.session.add(AllocationHistory(
        mentee_id=mentee.id,
        old_mentor_id=old_mentor,
        new_mentor_id=mentor.id,
        changed_by_id=current_user.id,
    ))
    db.session.commit()
    Notification.add(mentee.id, "mentor_assigned", f"Your mentor has been assigned to {mentor.full_name}.")
    AuditLog.log(current_user.id, "mentor_assignment", f"Assigned {mentee.username} to mentor {mentor.username}.")
    flash("Mentor allocation updated.", "success")
    return redirect(url_for("admin.allocations"))


@admin_bp.route("/substitutions", methods=["GET", "POST"])
@role_required("admin")
def substitutions():
    form = SubstitutionForm()
    form.mentee_id.choices = [(mentee.id, mentee.full_name) for mentee in User.query.filter_by(role="mentee", is_active=True).all()]
    form.substitute_mentor_id.choices = [(mentor.id, mentor.full_name) for mentor in User.query.filter_by(role="mentor", is_active=True).all()]
    substitutions = Substitution.query.order_by(Substitution.start_date.desc()).all()
    if form.validate_on_submit():
        mentee = User.query.filter_by(
            id=form.mentee_id.data,
            role="mentee",
            is_active=True,
        ).first()
        selected_substitute = User.query.filter_by(
            id=form.substitute_mentor_id.data,
            role="mentor",
            is_active=True,
        ).first()
        overlapping = Substitution.query.filter(
            Substitution.mentee_id == form.mentee_id.data,
            Substitution.active.is_(True),
            Substitution.start_date <= form.end_date.data,
            Substitution.end_date >= form.start_date.data,
        ).first()
        error = None
        if not mentee:
            error = "Choose an active mentee."
        elif not mentee.mentor_id or not User.query.filter_by(
            id=mentee.mentor_id,
            role="mentor",
            is_active=True,
        ).first():
            error = "Assign the mentee to an active mentor before scheduling a substitution."
        elif not selected_substitute:
            error = "Choose an active mentor as the substitute."
        elif selected_substitute.id == mentee.mentor_id:
            error = "The substitute must be different from the mentee's current mentor."
        elif form.start_date.data < date.today():
            error = "The substitution cannot start in the past."
        elif form.end_date.data < form.start_date.data:
            error = "The end date must be on or after the start date."
        elif overlapping:
            error = "This mentee already has an active substitution during those dates."
        if error:
            flash(error, "danger")
            return render_template(
                "admin/substitutions.html",
                form=form,
                substitutions=substitutions,
                today=date.today(),
            )
        substitution = Substitution(
            mentee_id=mentee.id,
            mentor_id=mentee.mentor_id,
            substitute_mentor_id=selected_substitute.id,
            reason=form.reason.data.strip(),
            start_date=form.start_date.data,
            end_date=form.end_date.data,
            active=True,
        )
        db.session.add(substitution)
        db.session.commit()
        Notification.add(mentee.id, "substitution", "A mentor substitution has been assigned to you.")
        AuditLog.log(
            current_user.id,
            "admin_substitution_create",
            f"Created a substitution for mentee ID {mentee.id}.",
        )
        flash("Substitution saved.", "success")
        return redirect(url_for("admin.substitutions"))
    return render_template(
        "admin/substitutions.html",
        form=form,
        substitutions=substitutions,
        today=date.today(),
    )


@admin_bp.route("/stats")
@role_required("admin")
def stats():
    mentors = User.query.filter_by(role="mentor", is_active=True).all()
    mentees = User.query.filter_by(role="mentee", is_active=True).all()
    meetings = Meeting.query.all()
    return render_template("admin/stats.html", mentors=mentors, mentees=mentees, meetings=meetings)


@admin_bp.route("/concerns")
@role_required("admin")
def concerns():
    concerns = Concern.query.order_by(Concern.created_at.desc()).all()
    return render_template(
        "admin/concerns.html",
        concerns=concerns,
        review_form=ConcernReviewForm(),
    )


@admin_bp.route("/concerns/<int:concern_id>", methods=["POST"])
@role_required("admin")
def update_concern(concern_id):
    concern = Concern.query.get_or_404(concern_id)
    form = ConcernReviewForm()
    if form.validate_on_submit():
        concern.status = form.status.data
        if form.response.data and form.response.data.strip():
            append_concern_response(concern, form.response.data, current_user.full_name, current_user.role)
        db.session.commit()
        add_notification(
            concern.mentee_id,
            "concern_update",
            f"Your concern was updated: {concern.title}",
            concern.id,
        )
        AuditLog.log(current_user.id, "concern_update", f"Updated concern {concern.id}.")
        flash("Concern update saved.", "success")
    else:
        flash("Choose a valid status and provide a reply of up to 2,000 characters.", "danger")
    return redirect(url_for("admin.concerns"))


@admin_bp.route("/audit")
@role_required("admin")
def audit():
    logs = AuditLog.query.order_by(AuditLog.created_at.desc()).all()
    return render_template("admin/audit.html", logs=logs)
